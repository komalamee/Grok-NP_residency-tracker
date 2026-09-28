import openpyxl, html, subprocess, weasyprint
from openpyxl.utils import get_column_letter
from PIL import Image, ImageChops
wb=openpyxl.load_workbook('sample-travel-log.xlsx')
for ws in wb:
    n=ws.max_column
    widths=[ws.column_dimensions[get_column_letter(i)].width*7.8 for i in range(1,n+1)]
    total=sum(widths)+40
    cols='<col style="width:40px">'+''.join(f'<col style="width:{w:.0f}px">' for w in widths)
    rows='<tr><th class="g"></th>'+''.join(f'<th class="g">{get_column_letter(i)}</th>' for i in range(1,n+1))+'</tr>'
    rows+=f'<tr><td class="g">1</td><td colspan="{n}" class="note">{html.escape(ws["A1"].value)}</td></tr>'
    for r in range(2,ws.max_row+1):
        rows+=f'<tr><td class="g">{r}</td>'
        for c in range(1,n+1):
            cell=ws.cell(row=r,column=c); v=html.escape(str(cell.value or ''))
            if r==2:
                bg='#'+cell.fill.fgColor.rgb[-6:]; fg='#'+cell.font.color.rgb[-6:]
                rows+=f'<td class="h" style="background:{bg};color:{fg}">{v}</td>'
            else: rows+=f'<td>{v}</td>'
        rows+='</tr>'
    tabs=''.join(f'<span class="tab{" on" if s.title==ws.title else ""}">{s.title}</span>' for s in wb)
    doc=f'''<html><head><style>
@page {{ size:{total+40:.0f}px 900px; margin:16px; }}
body {{ font-family:"DejaVu Sans",Arial,sans-serif; font-size:12px; color:#222; }}
.top {{ display:flex; justify-content:space-between; align-items:center; margin-bottom:8px; }}
.title {{ font-size:16px; font-weight:bold; }}
.ex {{ background:#fff3cd; border:1px solid #c9a227; color:#7a5d00; font-weight:bold; font-size:11px; padding:3px 8px; letter-spacing:1px; }}
table {{ border-collapse:collapse; table-layout:fixed; width:{total:.0f}px; }}
td,th {{ border:1px solid #d0d4da; padding:4px 6px; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; height:18px; }}
.g {{ background:#f1f3f4; color:#666; font-weight:normal; text-align:center; font-size:11px; }}
.note {{ font-style:italic; color:#555; }}
td.h {{ font-weight:bold; border-bottom:2px solid #8a939c; }}
.tabs {{ margin-top:8px; }} .tab {{ display:inline-block; padding:4px 12px; border:1px solid #d0d4da; background:#f1f3f4; margin-right:4px; color:#555; }}
.tab.on {{ background:#fff; color:#1a4d8f; font-weight:bold; border-top:3px solid #1a4d8f; }}
.legend {{ color:#666; font-size:10px; margin-top:6px; }}
.sw {{ display:inline-block; width:10px; height:10px; border:1px solid #bbb; vertical-align:middle; margin:0 3px 0 8px; }}
</style></head><body>
<div class="top"><span class="title">Nomad Pro – Travel log</span><span class="ex">EXAMPLE DATA</span></div>
<table>{cols}{rows}</table>
<div class="tabs">{tabs}</div>
<div class="legend">Fictional example rows. Rows 1–2 frozen, no formulas.<span class="sw" style="background:#DCE6F2"></span>editable<span class="sw" style="background:#EFEFEF"></span>written by Nomad Pro</div>
</body></html>'''
    name=ws.title.lower()
    weasyprint.HTML(string=doc).write_pdf(f'/tmp/sample/{name}.pdf')
    subprocess.run(['pdftoppm','-png','-r','150','-singlefile',f'/tmp/sample/{name}.pdf',f'sample-{name}'],check=True)
    p=f'sample-{name}.png'; im=Image.open(p).convert('RGB')
    b=ImageChops.difference(im,Image.new('RGB',im.size,(255,255,255))).getbbox()
    im=im.crop((max(0,b[0]-20),max(0,b[1]-20),min(im.width,b[2]+20),min(im.height,b[3]+20))); im.save(p); print(p,im.size)
